"""Regression tests for two guided-flow bugs reported live:

1. A default guided run must build a NON-managed adapter, and selecting a
   managed agent without its IDs must yield a clean, user-facing validation
   error (ops.AgentConfigError) rather than an opaque ValueError crash.
2. The Business Requirement step accepts an uploaded document (pdf/docx/txt/md)
   and extracts its text, behind auth.
"""

from __future__ import annotations

import asyncio
import io

import pytest
from fastapi.testclient import TestClient

from ascore import ops
from ascore.adapters.anthropic_simple import AnthropicSimpleAgent
from ascore.documents import MAX_DOCUMENT_BYTES, DocumentError, extract_text
from ascore.registry.sqlite_store import Registry
from ascore.server import nodes
from ascore.server.app import create_app
from tests.test_e2e_pipeline import RoutingFakeClient

CFG = {
    "models": {"agent_default": "agent-model", "judge_strong": "j", "judge_light": "l"},
    "harness": {"timeout_seconds": 10, "max_parallel": 5,
                "transport_retries": 1, "max_steps": 10},
    "scoring": {"calibration_threshold": 0.8},
    "paths": {"review_dir": "review/"},
}


# -- Bug 1: managed adapter never reached by default; clean error otherwise ---

class TestManagedAdapterGuard:
    def test_default_agent_config_is_reference(self):
        """The Agent-Under-Test node default must be the built-in reference
        agent — never managed."""
        cfg = nodes.AgentConfig()
        assert cfg.variant == "reference"

    def test_default_guided_run_builds_non_managed_adapter(self):
        """Building an adapter from the default agent config yields the built-in
        reference adapter, not a managed one."""
        cfg = nodes.AgentConfig()
        adapter = ops.build_adapter(
            CFG, variant=cfg.variant, agent_id=cfg.agent_id,
            client=RoutingFakeClient())
        assert isinstance(adapter, AnthropicSimpleAgent)

    def test_managed_without_ids_is_clean_validation_error(self):
        """build_adapter surfaces a user-facing AgentConfigError (not a bare
        ValueError crash) when managed is selected without IDs."""
        with pytest.raises(ops.AgentConfigError) as exc:
            ops.build_adapter(CFG, variant="managed", agent_id="x")
        assert "managed" in str(exc.value).lower()
        # AgentConfigError is a ValueError subclass: callers catching ValueError
        # still work, but the message is the friendly user-facing one.
        assert isinstance(exc.value, ValueError)
        assert str(exc.value) == ops.MANAGED_UNAVAILABLE_MSG

    def test_agent_node_rejects_managed_without_ids(self):
        """The Agent node fails fast with the friendly message if a managed
        agent slips in without IDs (and we're not deploying one)."""
        ctx = nodes.NodeContext(
            cfg=CFG, reg=None, execution_id="e", node_id="agent",
            emit=lambda *a, **k: None,
            wait_for_approval=None, cancelled=asyncio.Event())
        cfg = nodes.AgentConfig(variant="managed")
        with pytest.raises(ops.AgentConfigError, match="managed"):
            asyncio.run(nodes._run_agent(ctx, cfg, {}))


# -- Bug 2: document text extraction --------------------------------------

class TestExtractText:
    def test_txt(self):
        assert extract_text("req.txt", b"hello world") == "hello world"

    def test_md(self):
        out = extract_text("req.md", b"# Title\n\nbody")
        assert "Title" in out and "body" in out

    def test_pdf_roundtrip(self):
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("helvetica", size=12)
        pdf.cell(0, 10, "Refund within thirty days")
        data = pdf.output()
        out = extract_text("req.pdf", bytes(data))
        assert "Refund" in out and "thirty" in out

    def test_docx_roundtrip(self):
        import docx
        d = docx.Document()
        d.add_paragraph("Agents must never delete production data.")
        buf = io.BytesIO()
        d.save(buf)
        out = extract_text("req.docx", buf.getvalue())
        assert "never delete production data" in out

    def test_unsupported_type(self):
        with pytest.raises(DocumentError, match="unsupported"):
            extract_text("req.exe", b"MZ\x90\x00")

    def test_oversize(self):
        with pytest.raises(DocumentError, match="too large"):
            extract_text("req.txt", b"x" * (MAX_DOCUMENT_BYTES + 1))

    def test_empty(self):
        with pytest.raises(DocumentError, match="empty"):
            extract_text("req.txt", b"")


# -- Bug 2: the /api/documents/extract endpoint ----------------------------

CONFIG = """\
models: {agent_default: a, judge_strong: j, judge_light: l}
harness: {timeout_seconds: 10, max_parallel: 5, transport_retries: 1, max_steps: 10}
scoring: {calibration_threshold: 0.8}
live: {sample_rate: 0.05, drift_threshold: 0.15, drift_window_runs: 50}
paths: {registry_db: %(db)s, review_dir: %(r)s, calibration_dir: %(c)s}
auth: {token: "adm", tokens: {vw: viewer, op: operator}}
"""


@pytest.fixture
def authed_client(tmp_path):
    cfg = tmp_path / "config.yaml"
    cfg.write_text(CONFIG % {"db": tmp_path / "a.db", "r": tmp_path / "r",
                             "c": tmp_path / "c"})
    with TestClient(create_app(str(cfg), registry=Registry(tmp_path / "a.db"))) as c:
        yield c


def _post(client, data: bytes, name: str, token: str | None):
    headers = {"Authorization": f"Bearer {token}"} if token else {}
    return client.post("/api/documents/extract", headers=headers,
                       files={"file": (name, data, "application/octet-stream")})


class TestExtractEndpoint:
    def test_extract_requires_auth(self, authed_client):
        assert _post(authed_client, b"hi", "r.txt", None).status_code == 401

    def test_extract_forbids_viewer(self, authed_client):
        assert _post(authed_client, b"hi", "r.txt", "vw").status_code == 403

    def test_extract_txt_ok(self, authed_client):
        r = _post(authed_client, b"refund policy is 30 days", "r.txt", "op")
        assert r.status_code == 200
        body = r.json()
        assert body["text"] == "refund policy is 30 days"
        assert body["filename"] == "r.txt" and body["chars"] == 24

    def test_extract_bad_type_is_422(self, authed_client):
        r = _post(authed_client, b"MZ", "r.exe", "op")
        assert r.status_code == 422
        assert "unsupported" in r.json()["detail"]
